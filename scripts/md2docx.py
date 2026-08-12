"""Сборка .docx из markdown-документа по шаблону-образцу.

ЗАЧЕМ. `docs/CONCEPT.md` — живой документ (версионируется, диффается, правится в репо),
но заказчику концепция уходит файлом Word. Чтобы .docx не был непрозрачным бинарником,
который никто не может пересобрать, он генерируется этим скриптом из того же markdown.

КАК УСТРОЕНО. Готовый .docx берётся как ШАБЛОН: из него наследуются шрифты, поля страницы,
колонтитулы и стили заголовков, а тело документа собирается заново. Палитра повторяет
оригинал версии 1.0 (тёмная бирюза заголовков, зебра таблиц, цветные врезки).

Поддерживается ровно то подмножество markdown, которым написан CONCEPT.md: заголовки `#`…`####`,
таблицы, врезки `>`, списки `- `, инлайн `**жирный**` / `код` / [ссылка](url).

ЗАПУСК (Windows, кириллица — обязательно UTF-8):
    PYTHONUTF8=1 .venv/Scripts/python.exe scripts/md2docx.py \\
        docs/CONCEPT.md docs/CONCEPT_v1.0_2026-07-03.docx docs/CONCEPT_v2.0_2026-08-12.docx

ПРОВЕРКА РЕЗУЛЬТАТА. После сборки убедиться, что в тексте не осталось сырой разметки
(`**`, backtick, `](`) и что все идентификаторы требований `FR-*` доехали, — разметка,
разорванная переносом строки, иначе утекает в документ как есть.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

# --- палитра оригинала -------------------------------------------------------
DARK = "0B3D4F"      # заголовки, заливка шапки таблиц
BODY = "1F2933"      # основной текст
TEAL = "0B7285"      # трек A
AMBER = "96690B"     # трек B, предупреждения
GREEN = "2B6A30"     # общее ядро
RED = "A33333"       # ограничения
WHITE = "FFFFFF"
FILL_ZEBRA = "F6F9FB"
FILL_TEAL = "E3F2F5"
FILL_AMBER = "FDF3DF"
FILL_GREEN = "EDF3F0"

FONT = "Calibri"
PAGE_W = 9906  # twips: A4 (11906) минус поля 1000+1000


def shade(el, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_borders(table, color="C7D3DA", sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def style_run(run, *, color=BODY, size=18, bold=False, italic=False, mono=False):
    """size — в полупунктах, как в OOXML (18 = 9 pt)."""
    run.font.name = "Consolas" if mono else FONT
    run.font.size = Pt(size / 2)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), "Consolas" if mono else FONT)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text, *, color=BODY, size=18, bold=False):
    """Разбор инлайновой разметки: **жирный**, `код`, [ссылка](url)."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            style_run(par.add_run(part[2:-2]), color=color, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            style_run(par.add_run(part[1:-1]), color=color, size=size - 1,
                      bold=bold, mono=True)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            style_run(par.add_run(m.group(1)), color=TEAL, size=size, bold=bold)
        else:
            style_run(par.add_run(part), color=color, size=size, bold=bold)


def spacing(par, before=0, after=60, line=276):
    pf = par.paragraph_format
    pf.space_before = Pt(before / 20)
    pf.space_after = Pt(after / 20)
    pf.line_spacing = line / 240


# --- разбор markdown ---------------------------------------------------------
def parse_blocks(md):
    lines = md.split("\n")
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
            if len(cells) == 1 and len(cells[0]) == 1:
                blocks.append(("callout", cells[0][0]))
            else:
                blocks.append(("table", cells))
            continue
        if ln.startswith(">"):
            buf = []
            while i < len(lines) and (lines[i].startswith(">") or
                                      (buf and lines[i].strip() == "" and
                                       i + 1 < len(lines) and lines[i + 1].startswith(">"))):
                buf.append(re.sub(r"^>\s?", "", lines[i]))
                i += 1
            blocks.append(("quote", "\n".join(buf).strip()))
            continue
        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            blocks.append((f"h{lvl}", ln.lstrip("# ").strip()))
            i += 1
            continue
        if ln.startswith("- ") or ln.startswith("* "):
            blocks.append(("li", ln[2:].strip()))
            i += 1
            continue
        if ln.strip() == "---":
            blocks.append(("hr", ""))
            i += 1
            continue
        if ln.strip():
            buf = [ln.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                    r"^(\||>|#|- |\* |---)", lines[i]):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(buf)))
            continue
        i += 1
    return blocks


# --- сборка документа --------------------------------------------------------
def build(md_path, template, out_path):
    md = open(md_path, encoding="utf-8").read()
    doc = Document(template)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    blocks = parse_blocks(md)
    first_h1 = True

    for kind, payload in blocks:
        if kind == "hr":
            continue

        if kind.startswith("h"):
            lvl = int(kind[1])
            if lvl == 1 and first_h1:
                first_h1 = False
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                spacing(p, before=0, after=120)
                add_runs(p, payload, color=DARK, size=44, bold=True)
                continue
            p = doc.add_paragraph(style=f"Heading{min(lvl, 6)}")
            spacing(p, before=280 if lvl <= 2 else 200, after=100)
            size = {1: 30, 2: 24, 3: 21, 4: 19}.get(lvl, 18)
            add_runs(p, payload, color=DARK, size=size, bold=True)
            continue

        if kind == "p":
            p = doc.add_paragraph()
            spacing(p)
            add_runs(p, payload)
            continue

        if kind == "li":
            p = doc.add_paragraph(style="ListParagraph")
            spacing(p, after=40)
            p.paragraph_format.left_indent = Pt(18)
            style_run(p.add_run("•  "), color=TEAL, size=18, bold=True)
            add_runs(p, payload)
            continue

        if kind in ("quote", "callout"):
            text = payload
            fill = FILL_TEAL
            if text.startswith("⚠") or "⚠" in text[:80]:
                fill = FILL_AMBER
            elif text.startswith("**Урок") or text.startswith("**Вывод"):
                fill = FILL_GREEN
            t = doc.add_table(rows=1, cols=1)
            t.autofit = False
            set_borders(t, color="D8E3E8", sz=2)
            cell = t.cell(0, 0)
            cell.width = Twips(PAGE_W)
            cell_shade(cell, fill)
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            # переносы строк в markdown — мягкие: склеиваем в логические абзацы,
            # иначе **жирный**, разорванный переносом, попадает в текст сырым
            chunks, buf = [], []
            for raw in text.split("\n"):
                s = raw.strip()
                if not s:
                    if buf:
                        chunks.append(" ".join(buf))
                        buf = []
                    continue
                if s.startswith(("- ", "* ")):
                    if buf:
                        chunks.append(" ".join(buf))
                        buf = []
                    chunks.append(s)
                    continue
                buf.append(s)
            if buf:
                chunks.append(" ".join(buf))
            for j, chunk in enumerate(chunks):
                p = cell.add_paragraph()
                spacing(p, before=0, after=40 if j else 40)
                if chunk.lstrip().startswith(("- ", "* ")):
                    p.paragraph_format.left_indent = Pt(12)
                    style_run(p.add_run("•  "), color=AMBER if fill == FILL_AMBER else TEAL,
                              size=18, bold=True)
                    chunk = chunk.lstrip()[2:]
                add_runs(p, chunk, color=AMBER if fill == FILL_AMBER else DARK, size=19)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if kind == "table":
            rows = payload
            ncols = max(len(r) for r in rows)
            rows = [r + [""] * (ncols - len(r)) for r in rows]
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.autofit = False
            set_borders(t)
            width = PAGE_W // ncols
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    cell = t.cell(ri, ci)
                    cell.width = Twips(width)
                    par = cell.paragraphs[0]
                    spacing(par, before=0, after=0, line=240)
                    if ri == 0:
                        cell_shade(cell, DARK)
                        add_runs(par, val, color=WHITE, size=17, bold=True)
                    else:
                        if ri % 2 == 0:
                            cell_shade(cell, FILL_ZEBRA)
                        col = BODY
                        if "ТРЕК A" in val or val.startswith("**FR-"):
                            col = TEAL
                        elif "ТРЕК B" in val:
                            col = AMBER
                        elif "ЯДРО" in val:
                            col = GREEN
                        elif val.startswith("⏳") or val.startswith("⛔"):
                            col = RED
                        add_runs(par, val, color=col, size=17)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

    doc.save(out_path)
    print(f"OK -> {out_path}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3])
